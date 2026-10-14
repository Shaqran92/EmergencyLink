"""Build the Data Scientist results deck (docs/ds-results-slides.pptx) from results.json + charts."""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "evaluation" / "charts"
R = json.loads((ROOT / "evaluation" / "results.json").read_text())

NAVY = RGBColor(0x0F, 0x17, 0x2A); BLUE = RGBColor(0x25, 0x63, 0xEB); GREEN = RGBColor(0x16, 0xA3, 0x4A)
AMBER = RGBColor(0xF5, 0x9E, 0x0B); RED = RGBColor(0xDC, 0x26, 0x26); MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5


def text(slide, x, y, w, h, s, size=16, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
        p.space_after = Pt(6)
    return tb


def bullets(slide, x, y, w, h, items, size=15, color=NAVY):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        head, _, rest = it.partition("|")
        r = p.add_run(); r.text = "▸ "; r.font.color.rgb = BLUE; r.font.size = Pt(size); r.font.name = FONT
        r = p.add_run(); r.text = head; r.font.bold = bool(rest); r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
        if rest:
            r = p.add_run(); r.text = " " + rest; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
        p.space_after = Pt(9)
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def picture(slide, path, x, y, max_w, max_h):
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Inches(x + (max_w - w) / 2), Inches(y + (max_h - h) / 2), Inches(w), Inches(h))


def header(slide, title, kicker, n):
    rect(slide, 0, 0, W, 0.12, BLUE)
    text(slide, 0.6, 0.35, 11.5, 0.4, kicker.upper(), size=12, bold=True, color=BLUE)
    text(slide, 0.6, 0.7, 12, 0.8, title, size=28, bold=True)
    text(slide, 0.6, H - 0.45, 9, 0.3, "EmergencyLink · OpenHack 26 · HT-03 · Data Scientist · fictional data, illustrative values",
         size=10, color=MUTED)
    text(slide, W - 1.2, H - 0.45, 0.6, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def stat(slide, x, y, w, big, label, color):
    rect(slide, x, y, w, 1.45, LIGHT)
    rect(slide, x, y, 0.08, 1.45, color)
    text(slide, x + 0.25, y + 0.12, w - 0.35, 0.7, big, size=30, bold=True, color=color)
    text(slide, x + 0.25, y + 0.8, w - 0.35, 0.6, label, size=12, color=NAVY)


# 1 ─ title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, NAVY)
rect(s, 0.6, 2.1, 0.12, 2.3, BLUE)
text(s, 0.95, 1.95, 11, 0.5, "EMERGENCYLINK · OPENHACK 26 · HT-03", size=14, bold=True, color=RGBColor(0x93, 0xC5, 0xFD))
text(s, 0.95, 2.45, 11.5, 1.2, "An explainable hospital recommendation engine", size=34, bold=True, color=WHITE)
text(s, 0.95, 3.55, 11, 0.9, "Rules-first ranking that treats unknown data as unknown, explains every point, and leaves the decision to people.",
     size=18, color=RGBColor(0xCB, 0xD5, 0xE1))
text(s, 0.95, 5.6, 11, 0.4, "Data Scientist deliverables: engine · fictional dataset · test report · scoring logic", size=14, color=RGBColor(0x94, 0xA3, 0xB8))
text(s, 0.95, 6.0, 11, 0.4, "All hospitals and cases are fictional. Scores rank operational suitability. They are not clinical probabilities.",
     size=12, color=RGBColor(0x94, 0xA3, 0xB8))

# 2 ─ how it decides
s = prs.slides.add_slide(BLANK)
header(s, "Hard constraints first, then scoring", "How it decides", 2)
steps = ["Validate input", "Hard constraints", "Score eligible only", "Explain", "Human review", "Hospital accepts"]
for i, st in enumerate(steps):
    x = 0.6 + i * 2.05
    col = BLUE if i < 4 else MUTED
    b = rect(s, x, 1.75, 1.85, 0.75, col, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    text(s, x + 0.25, 1.8, 1.45, 0.65, st, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 8.85, 2.55, 4, 0.4, "outside the engine: MERN workflow", size=11, color=MUTED, align=PP_ALIGN.CENTER)
cols = [("ELIGIBLE", GREEN, "Every requirement passes on current (fresh or stale) data",
         ["Ranked with a 0–100 match score", "Stale / unconfirmed reports flagged", "Factor-by-factor explanation"]),
        ("NEEDS CONFIRMATION", AMBER, "Nothing failed, but a required resource is unknown, expired or conflicting",
         ["No score: we never invent one", "Reason shown, e.g. 'ICU report expired 5 h ago'", "ETA shown so staff know who to call"]),
        ("EXCLUDED", RED, "A hard constraint failed",
         ["Closed / diverting, rejected referral", "Required service or resource missing", "A high score can never buy it back"])]
for i, (t, c, sub, items) in enumerate(cols):
    x = 0.6 + i * 4.1
    rect(s, x, 3.15, 3.85, 3.55, LIGHT)
    rect(s, x, 3.15, 3.85, 0.55, c)
    text(s, x + 0.2, 3.2, 3.5, 0.45, t, size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 3.8, 3.5, 0.8, sub, size=13, bold=True)
    bullets(s, x + 0.15, 4.65, 3.6, 2, items, size=13)

# 3 ─ score
s = prs.slides.add_slide(BLANK)
header(s, "Match score: four transparent factors", "Scoring logic", 3)
rect(s, 0.6, 1.6, 12.1, 0.75, NAVY)
text(s, 0.8, 1.62, 11.8, 0.7, "Score = 0.40 × Travel + 0.30 × Freshness + 0.20 × Capacity + 0.10 × Capability", size=20, bold=True,
     color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, 0.6, 2.65, 5.4, 4.3, [
    "Travel|scored on the upper bound of the ETA range. Live traffic, map routing or straight-line, each labelled High / Medium / Low.",
    "Freshness|resource-specific decay (beds 30/120 min, CT 120/480). −15 if not hospital-confirmed. Weakest resource counts.",
    "Capacity|headroom above what the case needs. Tightest resource counts.",
    "Capability|preferred extras, e.g. a ventilator. Unknown ≠ matched.",
    "Unknown ETA|factor dropped, score marked 'incomplete', ranked after every complete score.",
], size=13)
picture(s, CH / "scoring_curves.png", 6.2, 2.6, 6.6, 4.2)

# 4 ─ demo
d1 = R["demo"]["step1"]; d2 = R["demo"]["step2"]
s = prs.slides.add_slide(BLANK)
header(s, "Demo: closest is not always best-evidenced", "Explainable result", 4)
picture(s, CH / "demo_step1_breakdown.png", 0.5, 1.6, 7.6, 3.4)
rect(s, 8.4, 1.65, 4.35, 5.2, LIGHT)
text(s, 8.6, 1.75, 4, 0.4, "Fictional cardiac case EL-DEMO-001", size=14, bold=True, color=BLUE)
top, sec = d1["candidates"][0], d1["candidates"][1]
bullets(s, 8.55, 2.2, 4.1, 4.6, [
    f"#1 {top['name']}|{top['matchScoreLabel']}, 14–17 min, verified data.",
    f"#2 {sec['name']}|closer (8–11 min), but its ICU report is 1.6 h old and unconfirmed, so {sec['matchScoreLabel'].lower()}.",
    f"{d1['needsConfirmation'][0]['name']}|not scored: {d1['needsConfirmation'][0]['reasons'][0].lower()}.",
    f"{d1['excluded'][0]['name']}|excluded: {d1['excluded'][0]['reasons'][0].lower()}.",
], size=12.5)
rect(s, 0.6, 5.25, 7.5, 1.6, RGBColor(0xFE, 0xF2, 0xF2))
rect(s, 0.6, 5.25, 0.08, 1.6, RED)
text(s, 0.85, 5.32, 7.1, 0.4, "Failure path: Crescent Heart rejects the referral", size=14, bold=True, color=RED)
text(s, 0.85, 5.75, 7.1, 1.1, [f"Re-rank → only {d2['candidates'][0]['name']} remains, still flagged Stale + Unconfirmed.",
                               "Engine warns: “Only one eligible candidate” · “Top candidate carries data-quality flags”."], size=12.5)

# 5 ─ network
s = prs.slides.add_slide(BLANK)
header(s, "10 cases × 12 hospitals: uncertainty stays visible", "Network view", 5)
picture(s, CH / "case_outcomes.png", 0.5, 1.55, 7.9, 5.3)
bullets(s, 8.7, 1.8, 4.1, 5, [
    "Stroke (EL-0003)|1 eligible. Riverside Neuro is held back by conflicting CT reports from two staff members.",
    "Burns (EL-0004)|only Sunrise qualifies. Harborline reports its burn beds full.",
    "No caller location (EL-0008)|all scores marked 'incomplete', never inflated.",
    "Dialysis (EL-0010)|Saffron Valley has no map location, so it ranks after Harborline despite a higher partial score.",
    "Data quality|a future-dated record (device clock error) was detected and ignored.",
], size=12)

# 6 ─ testing
inv = R["invariants"]; sc = R["scenarios"]; lat = {x["hospitals"]: x for x in R["latency"]}
npass = sum(x["passed"] for x in sc)
s = prs.slides.add_slide(BLANK)
header(s, "Tested against expectations written before running", "Evaluation", 6)
tests_n = R["pytest"]["summary"].split(" passed")[0]
stat(s, 0.6, 1.65, 2.95, f"{tests_n}", "automated tests passing (pytest)", GREEN)
stat(s, 3.7, 1.65, 2.95, f"{npass}/{len(sc)}", "predefined scenarios passed", GREEN)
stat(s, 6.8, 1.65, 2.95, f"{inv['totalViolations']}", f"safety violations in {inv['requests']:,} random requests", GREEN)
stat(s, 9.9, 1.65, 2.85, f"{lat[200]['p95ms']:.0f} ms", "p95 latency, 200 hospitals", BLUE)
bullets(s, 0.6, 3.45, 6.0, 3.6, [
    "Scenarios|stale, expired, unknown, conflicting, future-dated, closed, rejected, insufficient count, tie-breaks, ETA-unknown…",
    "Invariants|no ranked hospital ever fails a constraint; contributions always sum to the score; complete scores outrank incomplete ones.",
    "Determinism|identical input gives identical output, independent of input order.",
    "Not covered here|offline sync, authorisation, end-to-end workflow. Those are owned by the MERN and Frontend members.",
], size=13)
picture(s, CH / "latency.png", 6.9, 3.4, 5.9, 3.5)

# 7 ─ robustness
s = prs.slides.add_slide(BLANK)
header(s, "Weights are provisional, but the top pick is stable", "Robustness & data quality", 7)
picture(s, CH / "weight_sensitivity.png", 0.4, 1.6, 6.4, 3.9)
picture(s, CH / "resource_freshness.png", 6.9, 1.6, 6.0, 3.9)
flips = sum(x["cases"] - x["top1Unchanged"] for x in R["sensitivity"])
taus = [x["meanKendallTau"] for x in R["sensitivity"] if x["meanKendallTau"] is not None]
text(s, 0.6, 5.65, 6.1, 1.3, [f"±0.05 / ±0.10 on each weight: top-1 changed in {flips} of {len(R['sensitivity']) * R['sensitivity'][0]['cases']} runs; "
                              f"rank agreement τ ≥ {min(taus):.2f}.", "Stable ≠ correct: the weights still need expert review."], size=13)
q = R["datasetQuality"]
text(s, 7.0, 5.65, 5.8, 1.3, [f"{q['freshSharePct']}% of latest reports fresh · {q['verifiedSharePct']}% hospital-confirmed.",
                              "Expired reports are shown as Unknown, never as a stale 'Available'."], size=13)

# 8 ─ stretch
fc = R["forecast"]; sim = R["surge"]
s = prs.slides.add_slide(BLANK)
header(s, "Stretch: surge simulation & demand forecast", "Optional extras", 8)
picture(s, CH / "surge_simulation.png", 0.4, 1.55, 6.4, 3.6)
picture(s, CH / "forecast_synthetic.png", 6.9, 1.55, 6.0, 3.6)
text(s, 0.6, 5.35, 6.1, 1.6, [f"{sim['casesSimulated']} trauma cases in 60 min, with a trauma-centre closure and a CT outage.",
                              f"Capacity runs out at minute 30: {sim['unplacedCases']} cases get an explicit 'no eligible hospital' warning instead of a made-up recommendation."], size=12.5)
gb = fc["metrics"]["Gradient boosting (Poisson)"]["MAE"]; bb = fc["metrics"][fc["bestBaseline"]]["MAE"]
text(s, 7.0, 5.35, 5.8, 1.6, [f"24-h-ahead arrivals: gradient boosting MAE {gb:.2f} vs best baseline {bb:.2f} ({fc['mlImprovementOverBestBaselinePct']:+.1f}%).",
                              "Synthetic data only. This demonstrates the method, not real accuracy."], size=12.5)

# 9 ─ safety & limits
s = prs.slides.add_slide(BLANK)
header(s, "What we claim, and what we don't", "Safety & limitations", 9)
rect(s, 0.6, 1.65, 5.9, 5.2, RGBColor(0xF0, 0xFD, 0xF4)); rect(s, 0.6, 1.65, 5.9, 0.55, GREEN)
text(s, 0.8, 1.7, 5.5, 0.45, "The engine does", size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, 0.75, 2.35, 5.6, 4.4, ["Check hard requirements before any scoring",
                                  "Show source, time and confirmation for every value",
                                  "Say 'Unknown' / 'Insufficient data' instead of guessing",
                                  "Explain every point of every score",
                                  "Log input hash + score version for audit",
                                  "Require human review and explicit hospital acceptance"], size=14)
rect(s, 6.85, 1.65, 5.9, 5.2, RGBColor(0xFE, 0xF2, 0xF2)); rect(s, 6.85, 1.65, 5.9, 0.55, RED)
text(s, 7.05, 1.7, 5.5, 0.45, "The engine does not", size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, 7.0, 2.35, 5.6, 4.4, ["Diagnose, triage or set clinical priority",
                                 "Confirm a destination or dispatch anyone",
                                 "Predict acceptance or patient outcomes",
                                 "Claim validated weights: all are design assumptions",
                                 "Turn synthetic results into real-world evidence"], size=14)

out = ROOT / "docs" / "ds-results-slides.pptx"
prs.save(out)
print("saved", out)
